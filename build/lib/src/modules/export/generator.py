import gc
import hashlib
import logging
import os
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Optional

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from src.core.config import settings
from src.core.logger import logger

logger_export = logging.getLogger("meeyg.export")

WP_COLUMNS = [
    "post_title",
    "sku",
    "regular_price",
    "sale_price",
    "stock",
    "stock_status",
    "categories",
    "tags",
    "short_description",
    "description",
    "images",
    "attributes",
    "manage_stock",
    "backorders",
    "tax_status",
    "meta",
]


@dataclass
class ExportConfig:
    supplier_ids: Optional[list[int]] = None
    category_ids: Optional[list[int]] = None
    ready_only: bool = False
    available_only: bool = False
    format: str = "xlsx"
    encoding: str = "utf-8-sig"
    category_separator: str = ">"
    image_separator: str = "|"
    attribute_format: str = "name:value"
    output_dir: Optional[Path] = None
    chunk_size: int = 1000

    def __post_init__(self):
        if self.output_dir is None:
            self.output_dir = settings.data_dir / "exports"
        self.output_dir.mkdir(parents=True, exist_ok=True)


@dataclass
class ExportStats:
    total_products: int = 0
    exported_products: int = 0
    duplicate_skus: int = 0
    empty_titles: int = 0
    empty_prices: int = 0
    elapsed: float = 0.0
    file_size: int = 0
    output_path: str = ""


class ExportEngine:
    def __init__(
        self,
        config: ExportConfig,
        log_callback: Optional[Callable[[str], None]] = None,
        progress_callback: Optional[Callable[[int, int, str], None]] = None,
    ):
        self.config = config
        self.log_callback = log_callback or (lambda m: logger_export.info(m))
        self.progress_callback = progress_callback or (lambda p, t, m: None)
        self._cancelled = False
        self.stats = ExportStats()

    def cancel(self) -> None:
        self._cancelled = True

    def _log(self, msg: str) -> None:
        self.log_callback(msg)

    def _progress(self, pct: int, total: int, msg: str) -> None:
        self.progress_callback(pct, total, msg)

    def _fetch_products(self, db_session) -> list[dict]:
        from src.database.models import Category, Product, ProductAttribute, Supplier

        self._log("Fetching products from database...")

        query = db_session.query(
            Product.id,
            Product.external_sku,
            Product.title,
            Product.description,
            Product.price,
            Product.currency,
            Product.is_available,
            Product.is_ready_for_export,
            Product.image_urls,
            Product.created_at,
            Supplier.name.label("supplier_name"),
            Category.name.label("category_name"),
        ).outerjoin(Supplier, Product.supplier_id == Supplier.id).outerjoin(
            Category, Product.category_id == Category.id
        )

        if self.config.supplier_ids:
            query = query.filter(Product.supplier_id.in_(self.config.supplier_ids))
        if self.config.category_ids:
            query = query.filter(Product.category_id.in_(self.config.category_ids))
        if self.config.ready_only:
            query = query.filter(Product.is_ready_for_export == True)
        if self.config.available_only:
            query = query.filter(Product.is_available == True)

        products = query.order_by(Product.id).all()
        self.stats.total_products = len(products)
        self._log(f"Found {len(products)} products matching filters")

        attr_query = db_session.query(
            ProductAttribute.product_id,
            ProductAttribute.name,
            ProductAttribute.value,
        )
        if self.config.supplier_ids:
            attr_query = attr_query.join(Product).filter(
                Product.supplier_id.in_(self.config.supplier_ids)
            )
        if self.config.category_ids:
            attr_query = attr_query.join(Product).filter(
                Product.category_id.in_(self.config.category_ids)
            )

        attributes_rows = attr_query.all()
        attr_map: dict[int, list[tuple[str, str]]] = {}
        for row in attributes_rows:
            attr_map.setdefault(row.product_id, []).append((row.name, row.value))

        for p in products:
            p_attrs = attr_map.get(p.id, [])
            if self.config.attribute_format == "name:value":
                attr_str = "; ".join(f"{n}: {v}" for n, v in p_attrs)
            else:
                attr_str = "; ".join(f"{n}={v}" for n, v in p_attrs)

            image_urls = []
            if p.image_urls:
                import json
                try:
                    image_urls = json.loads(p.image_urls)
                except (json.JSONDecodeError, TypeError):
                    pass

            yield {
                "id": p.id,
                "external_sku": p.external_sku,
                "title": p.title,
                "description": p.description or "",
                "price": p.price,
                "currency": p.currency,
                "is_available": p.is_available,
                "supplier_name": p.supplier_name,
                "category_name": p.category_name,
                "attributes": attr_str,
                "image_urls": image_urls,
                "created_at": p.created_at,
            }

    def _transform_to_wp_format(self, products: list[dict]) -> list[dict]:
        self._log("Transforming data to WP All Import format...")
        rows = []
        seen_skus: set[str] = set()

        for i, p in enumerate(products):
            if self._cancelled:
                break

            sku = p["external_sku"] or f"MEYG-{p['id']}"
            if sku in seen_skus:
                base_sku = sku
                counter = 1
                while sku in seen_skus:
                    sku = f"{base_sku}-{counter}"
                    counter += 1
                self.stats.duplicate_skus += 1
            seen_skus.add(sku)

            title = (p["title"] or "").strip()
            if not title:
                title = f"Product {p['id']}"
                self.stats.empty_titles += 1

            price = p["price"]
            if price is None:
                self.stats.empty_prices += 1

            stock = 1 if p["is_available"] else 0
            stock_status = "instock" if p["is_available"] else "outofstock"

            category = p["category_name"] or "Uncategorized"

            images_str = self.config.image_separator.join(p["image_urls"]) if p["image_urls"] else ""

            short_desc = ""
            if p["description"]:
                short_desc = (p["description"][:300] + "...") if len(p["description"]) > 300 else p["description"]

            meta_fields = {
                "supplier": p["supplier_name"],
                "currency": p["currency"] or "",
                "imported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "source_id": str(p["id"]),
            }

            rows.append({
                "post_title": title,
                "sku": sku,
                "regular_price": f"{price:.2f}" if price is not None else "",
                "sale_price": "",
                "stock": str(stock),
                "stock_status": stock_status,
                "categories": category,
                "tags": "",
                "short_description": short_desc,
                "description": p["description"] or "",
                "images": images_str,
                "attributes": p["attributes"],
                "manage_stock": "yes" if p["is_available"] else "no",
                "backorders": "no",
                "tax_status": "taxable",
                "meta": "; ".join(f"{k}: {v}" for k, v in meta_fields.items()),
            })

            if (i + 1) % 100 == 0:
                pct = int(((i + 1) / max(self.stats.total_products, 1)) * 100)
                self._progress(pct, self.stats.total_products, f"Transformed {i + 1} products")

        self.stats.exported_products = len(rows)
        return rows

    def _export_to_excel(self, rows: list[dict]) -> str:
        self._log("Generating Excel file with openpyxl...")

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        supplier_name = "all"
        if self.config.supplier_ids and len(self.config.supplier_ids) == 1:
            supplier_name = f"supplier_{self.config.supplier_ids[0]}"
        filename = f"{supplier_name}_{timestamp}.xlsx"
        output_path = self.config.output_dir / filename

        df = pd.DataFrame(rows, columns=WP_COLUMNS)

        wb = Workbook()
        ws = wb.active
        ws.title = "Products"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, col_name in enumerate(WP_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        price_cols = {"regular_price", "sale_price"}
        date_cols = set()

        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, col_name in enumerate(WP_COLUMNS, 1):
                value = row_data.get(col_name, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)

                if col_name in price_cols and value:
                    try:
                        cell.number_format = '#,##0.00'
                    except ValueError:
                        pass

                if row_idx % 2 == 0:
                    cell.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")

        ws.auto_filter.ref = f"A1:{get_column_letter(len(WP_COLUMNS))}{len(rows) + 1}"

        for col_idx in range(1, len(WP_COLUMNS) + 1):
            max_len = 12
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
                if row[0]:
                    max_len = max(max_len, min(len(str(row[0])), 50))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 2

        ws.freeze_panes = "B2"

        wb.save(str(output_path))

        file_size = os.path.getsize(output_path)
        self.stats.file_size = file_size
        self.stats.output_path = str(output_path)

        self._log(f"Excel saved: {output_path} ({file_size / 1024:.1f} KB)")
        return str(output_path)

    def _export_to_csv(self, rows: list[dict]) -> str:
        self._log("Generating CSV file...")

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        supplier_name = "all"
        if self.config.supplier_ids and len(self.config.supplier_ids) == 1:
            supplier_name = f"supplier_{self.config.supplier_ids[0]}"
        filename = f"{supplier_name}_{timestamp}.csv"
        output_path = self.config.output_dir / filename

        df = pd.DataFrame(rows, columns=WP_COLUMNS)
        df.to_csv(
            str(output_path),
            index=False,
            encoding=self.config.encoding,
            sep=";",
            quoting=1,
        )

        file_size = os.path.getsize(output_path)
        self.stats.file_size = file_size
        self.stats.output_path = str(output_path)

        self._log(f"CSV saved: {output_path} ({file_size / 1024:.1f} KB)")
        return str(output_path)

    def run(self, db_session) -> dict[str, Any]:
        self._cancelled = False
        self.stats = ExportStats()
        start_time = time.time()

        self._log(f"Export started (format={self.config.format})")

        products = list(self._fetch_products(db_session))
        if not products:
            self._log("No products found matching filters")
            return {"success": False, "error": "No products found", "stats": self.stats}

        self._progress(10, self.stats.total_products, "Fetching complete")

        if self._cancelled:
            return {"success": False, "error": "cancelled", "stats": self.stats}

        rows = self._transform_to_wp_format(products)
        if not rows:
            return {"success": False, "error": "No rows after transformation", "stats": self.stats}

        self._progress(60, self.stats.total_products, "Transformation complete")

        if self._cancelled:
            return {"success": False, "error": "cancelled", "stats": self.stats}

        if self.config.format == "csv":
            output_path = self._export_to_csv(rows)
        else:
            output_path = self._export_to_excel(rows)

        self.stats.elapsed = time.time() - start_time
        self._progress(100, self.stats.total_products, "Export complete")

        del products
        del rows
        gc.collect()

        self._log(f"Export complete in {self.stats.elapsed:.1f}s")

        return {
            "success": True,
            "output_path": output_path,
            "stats": self.stats,
        }

    def validate_before_export(self, db_session) -> list[str]:
        from src.database.models import Product

        warnings = []

        query = db_session.query(Product)
        if self.config.supplier_ids:
            query = query.filter(Product.supplier_id.in_(self.config.supplier_ids))
        if self.config.category_ids:
            query = query.filter(Product.category_id.in_(self.config.category_ids))
        if self.config.ready_only:
            query = query.filter(Product.is_ready_for_export == True)

        products = query.all()

        if not products:
            warnings.append("No products match the current filters")
            return warnings

        empty_titles = sum(1 for p in products if not p.title or not p.title.strip())
        if empty_titles:
            warnings.append(f"{empty_titles} products have empty titles")

        empty_prices = sum(1 for p in products if p.price is None)
        if empty_prices:
            warnings.append(f"{empty_prices} products have no price")

        empty_skus = sum(1 for p in products if not p.external_sku)
        if empty_skus:
            warnings.append(f"{empty_skus} products have no SKU (will be auto-generated)")

        return warnings
